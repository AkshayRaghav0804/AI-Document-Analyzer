import io

import pandas as pd
from docx import Document

from docanalyzer.extractors import build_corpus, extract_file


def _docx_bytes() -> bytes:
    doc = Document()
    doc.add_paragraph("Quarterly revenue was 4.2 million.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Region"
    table.rows[0].cells[1].text = "North"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_docx_text_and_tables():
    result = extract_file("report.docx", _docx_bytes())
    assert result.ok
    assert "4.2 million" in result.text
    assert "[Table 1]" in result.text and "North" in result.text


def test_small_excel_is_dumped_in_full():
    buf = io.BytesIO()
    pd.DataFrame({"a": [1, 2], "b": ["x", "y"]}).to_excel(buf, index=False)
    result = extract_file("small.xlsx", buf.getvalue())
    assert result.ok and not result.large_table
    assert "x" in result.text
    assert len(result.dataframes) == 1


def test_large_excel_is_summarised():
    buf = io.BytesIO()
    pd.DataFrame({"n": range(500), "c": ["a", "b"] * 250}).to_excel(buf, index=False)
    result = extract_file("big.xlsx", buf.getvalue())
    assert result.large_table
    assert "Numeric statistics" in result.text
    assert "500 rows" in result.text


def test_csv_supported():
    result = extract_file("data.csv", b"a,b\n1,2\n3,4\n")
    assert result.ok and result.kind == "csv"
    assert list(result.dataframes["data"].columns) == ["a", "b"]


def test_unsupported_type_fails_softly():
    result = extract_file("notes.txt", b"hello")
    assert not result.ok and "Unsupported" in result.note


def test_corrupt_pdf_fails_softly():
    result = extract_file("broken.pdf", b"not a pdf")
    assert not result.ok and result.note


def test_build_corpus_skips_failed_files():
    ok = extract_file("data.csv", b"a\n1\n")
    bad = extract_file("x.txt", b"")
    corpus = build_corpus([ok, bad])
    assert "data.csv" in corpus and "x.txt" not in corpus
