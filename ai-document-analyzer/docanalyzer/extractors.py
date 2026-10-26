"""Turn uploaded files into text (and DataFrames for spreadsheets).

Nothing here imports Streamlit, so it can be unit-tested and reused anywhere.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd

from . import config

try:
    import pymupdf as fitz  # PyMuPDF (new import name)

    HAS_PYMUPDF = True
except ImportError:
    try:
        import fitz  # older PyMuPDF releases

        HAS_PYMUPDF = True
    except ImportError:  # pragma: no cover
        HAS_PYMUPDF = False

try:
    import pdfplumber

    HAS_PDFPLUMBER = True
except ImportError:  # pragma: no cover
    HAS_PDFPLUMBER = False


@dataclass
class ExtractedDocument:
    name: str
    text: str
    kind: str  # pdf | docx | excel | csv | unsupported
    ok: bool = True
    note: str = ""
    large_table: bool = False
    dataframes: dict[str, pd.DataFrame] = field(default_factory=dict)

    @property
    def chars(self) -> int:
        return len(self.text)


# --------------------------------------------------------------------------- PDF
def _pdf_pymupdf(data: bytes) -> str:
    if not HAS_PYMUPDF:
        return ""
    pages = []
    with fitz.open(stream=data, filetype="pdf") as pdf:
        for i, page in enumerate(pdf, 1):
            text = page.get_text("text")
            if text.strip():
                pages.append(f"[Page {i}]\n{text}")
    return "\n\n".join(pages)


def _pdf_pdfplumber(data: bytes) -> str:
    if not HAS_PDFPLUMBER:
        return ""
    pages = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for i, page in enumerate(pdf.pages, 1):
            text = page.extract_text()
            if text:
                pages.append(f"[Page {i}]\n{text}")
    return "\n\n".join(pages)


def extract_pdf(data: bytes) -> str:
    for extractor in (_pdf_pymupdf, _pdf_pdfplumber):
        try:
            text = extractor(data)
        except Exception:
            continue
        if len(text.strip()) > 100:
            return text
    raise ValueError(
        "No readable text found. The PDF may be scanned (image only); run OCR first."
    )


# -------------------------------------------------------------------------- Word
def extract_docx(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for n, table in enumerate(doc.tables, 1):
        parts.append(f"\n[Table {n}]")
        for row in table.rows:
            parts.append("\t".join(cell.text.strip() for cell in row.cells))
    if not parts:
        raise ValueError("The document contains no text.")
    return "\n".join(parts)


# ------------------------------------------------------------------- Spreadsheet
def summarise_dataframe(df: pd.DataFrame, sheet: str) -> str:
    out = [f"Sheet: {sheet} | {len(df):,} rows x {len(df.columns)} columns"]
    cols = list(df.columns[:50])
    out.append("Columns: " + ", ".join(f'"{c}"' for c in cols))
    if len(df.columns) > 50:
        out.append(f"... and {len(df.columns) - 50} more columns")
    out.append("\nFirst 10 rows:\n" + df.head(10).to_string(max_colwidth=40))
    out.append("\nLast 5 rows:\n" + df.tail(5).to_string(max_colwidth=40))

    numeric = df.select_dtypes(include="number")
    if not numeric.empty:
        out.append("\nNumeric statistics:\n" + numeric.describe().to_string(max_cols=20))

    categorical = df.select_dtypes(include=["object", "string", "category"])
    if not categorical.empty:
        lines = []
        for col in categorical.columns[:15]:
            unique = df[col].nunique()
            line = f"- {col}: {unique:,} unique"
            if unique <= 20:
                top = df[col].value_counts().head(5).index.tolist()
                line += f" (top: {', '.join(map(str, top))})"
            lines.append(line)
        out.append("\nCategorical columns:\n" + "\n".join(lines))
    return "\n".join(out)


def extract_spreadsheet(data: bytes, name: str, kind: str) -> ExtractedDocument:
    if kind == "csv":
        frames = {"data": pd.read_csv(io.BytesIO(data))}
    else:
        frames = pd.read_excel(io.BytesIO(data), sheet_name=None)

    parts, large = [], False
    for sheet, df in frames.items():
        header = f"--- {name} | sheet: {sheet} ---"
        if len(df) > config.EXCEL_FULL_DUMP_ROWS or len(df.columns) > config.EXCEL_FULL_DUMP_COLS:
            large = True
            parts.append(header + "\n" + summarise_dataframe(df, sheet))
        else:
            parts.append(header + "\n" + df.to_string(max_colwidth=50))

    total_rows = sum(len(d) for d in frames.values())
    return ExtractedDocument(
        name=name,
        text="\n\n".join(parts),
        kind="csv" if kind == "csv" else "excel",
        large_table=large,
        dataframes=frames,
        note=f"{len(frames)} sheet(s), {total_rows:,} rows",
    )


# ------------------------------------------------------------------------ Router
def extract_file(name: str, data: bytes) -> ExtractedDocument:
    """Extract one file. Never raises: failures come back with ok=False."""
    ext = Path(name).suffix.lower().lstrip(".")
    try:
        if ext == "pdf":
            text = extract_pdf(data)
            return ExtractedDocument(name, text, "pdf", note=f"{text.count('[Page ')} page(s)")
        if ext == "docx":
            return ExtractedDocument(name, extract_docx(data), "docx")
        if ext in ("xlsx", "xls"):
            return extract_spreadsheet(data, name, "excel")
        if ext == "csv":
            return extract_spreadsheet(data, name, "csv")
        return ExtractedDocument(name, "", "unsupported", ok=False, note=f"Unsupported type: .{ext}")
    except Exception as exc:
        return ExtractedDocument(name, "", ext or "unknown", ok=False, note=str(exc))


def build_corpus(docs: list[ExtractedDocument]) -> str:
    """Join all successfully extracted documents into one prompt-ready string."""
    return "\n\n".join(
        f"=== FILE: {d.name} ===\n{d.text}" for d in docs if d.ok
    )
